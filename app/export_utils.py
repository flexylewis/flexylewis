"""Export utilities for scraped datasets."""
from __future__ import annotations

import csv
import importlib.util
from pathlib import Path
from typing import Iterable, List

from app.models import ScrapeResult


def export_csv(path: Path, rows: Iterable[ScrapeResult]) -> None:
    """Export results to a CSV file with UTF-8 encoding."""

    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["URL", "Title", "Content", "Price", "Links"])
        for row in rows:
            writer.writerow([row.url, row.title, row.content, row.price, row.links_text()])


def export_xlsx(path: Path, rows: Iterable[ScrapeResult]) -> None:
    """Export results to an XLSX file using openpyxl."""

    if importlib.util.find_spec("openpyxl") is None:
        raise RuntimeError("openpyxl is required for XLSX export.")

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["URL", "Title", "Content", "Price", "Links"])
    for row in rows:
        sheet.append([row.url, row.title, row.content, row.price, row.links_text()])
    workbook.save(path)


def export_docx(path: Path, rows: Iterable[ScrapeResult]) -> None:
    """Export results to a DOCX document using python-docx."""

    if importlib.util.find_spec("docx") is None:
        raise RuntimeError("python-docx is required for DOCX export.")

    from docx import Document

    document = Document()
    for index, row in enumerate(rows, start=1):
        document.add_heading(f"Record {index}", level=2)
        document.add_paragraph(f"URL: {row.url}")
        document.add_paragraph(f"Title: {row.title}")
        document.add_paragraph(f"Content: {row.content}")
        document.add_paragraph(f"Price: {row.price}")
        document.add_paragraph(f"Links: {row.links_text()}")
    document.save(path)


def export_html(path: Path, rows: Iterable[ScrapeResult]) -> None:
    """Export results to a basic HTML document."""

    lines: List[str] = ["<html><body>"]
    for index, row in enumerate(rows, start=1):
        lines.append(f"<h2>Record {index}</h2>")
        lines.append(f"<p><strong>URL:</strong> {row.url}</p>")
        lines.append(f"<p><strong>Title:</strong> {row.title}</p>")
        lines.append(f"<p><strong>Content:</strong> {row.content}</p>")
        lines.append(f"<p><strong>Price:</strong> {row.price}</p>")
        lines.append(f"<p><strong>Links:</strong> {row.links_text()}</p>")
    lines.append("</body></html>")

    path.write_text("\n".join(lines), encoding="utf-8")


def export_text(path: Path, rows: Iterable[ScrapeResult]) -> None:
    """Export results to a plain text file."""

    blocks: List[str] = []
    for row in rows:
        blocks.append(
            "\n".join(
                [
                    f"URL: {row.url}",
                    f"Title: {row.title}",
                    f"Content: {row.content}",
                    f"Price: {row.price}",
                    f"Links: {row.links_text()}",
                ]
            )
        )
    path.write_text("\n\n".join(blocks), encoding="utf-8")
